"""Lectura de listas de precios de proveedores (xlsx / csv).

Funciones puras: devuelven [(numero_de_fila, {encabezado_normalizado: valor})],
salteando filas vacías. Los encabezados se normalizan para comparar contra el
perfil sin importar mayúsculas, acentos ni espacios.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata

import openpyxl

Row = tuple[int, dict[str, object]]


def normalize_header(value) -> str:
    if value is None:
        return ''
    text = unicodedata.normalize('NFKD', str(value))
    text = ''.join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', text).strip().lower()


def parse_number(value, decimal: str) -> float | None:
    """Convierte un valor de planilla a float. `decimal` es ',' o '.'.

    Números nativos se devuelven tal cual. Texto: se quitan símbolos de moneda
    y espacios, se elimina el separador de miles y se normaliza el decimal.
    Devuelve None si no es un número.
    """
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r'[^\d,.\-]', '', str(value))
    if not text:
        return None
    thousands = '.' if decimal == ',' else ','
    text = text.replace(thousands, '').replace(decimal, '.')
    try:
        return float(text)
    except ValueError:
        return None


def _rows_from_matrix(matrix, header_row: int) -> list[Row]:
    if len(matrix) < header_row:
        raise ValueError(f'El archivo tiene menos de {header_row} filas.')
    headers = [normalize_header(h) for h in matrix[header_row - 1]]
    rows: list[Row] = []
    for offset, raw in enumerate(matrix[header_row:], start=header_row + 1):
        if not any(cell not in (None, '') for cell in raw):
            continue
        rows.append((offset, {
            header: raw[idx] if idx < len(raw) else None
            for idx, header in enumerate(headers) if header
        }))
    return rows


def read_xlsx(content: bytes, sheet_name: str | bool, header_row: int) -> list[Row]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f'La hoja "{sheet_name}" no existe. Hojas: {", ".join(workbook.sheetnames)}.'
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.worksheets[0]
    matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    return _rows_from_matrix(matrix, header_row)


def read_csv(content: bytes, delimiter: str, encoding: str, header_row: int) -> list[Row]:
    text = content.decode(encoding)
    matrix = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    # csv.reader devuelve [] para líneas vacías: se mantienen para no correr la numeración
    matrix = [row if row else [] for row in matrix]
    return _rows_from_matrix(matrix, header_row)
